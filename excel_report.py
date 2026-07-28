from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Placeholders until the register actually varies by shift -- every
# register processed so far runs the same single shift, so there's
# nothing yet to read this from.
_SHIFT_CODE = "G"
_SHIFT_TIME = "09:00-18:00"

_COLUMNS = [
    "Code", "Name", "Duty Date", "Shift Code", "ShiftTime",
    "In Time", "Out Time", "Status", "Remark",
]
_COLUMN_WIDTHS = (10, 20, 12, 10, 14, 10, 10, 10, 20)

# cell_highlight/validation flag names (see validation_engine.py) mapped
# to the column they land on in this layout -- Code/Name/Shift Code/
# ShiftTime aren't extracted per-record so they never carry a flag.
_FIELD_TO_COLUMN = {
    "date": "Duty Date",
    "in_time": "In Time",
    "out_time": "Out Time",
    "status": "Status",
    "remark": "Remark",
}

_HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_INVALID_FILL = PatternFill(start_color="FFB3B3", end_color="FFB3B3", fill_type="solid")
_CENTER = Alignment(horizontal="center")

_THIN_SIDE = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)

# Status-specific tint on just the Status cell, purely for at-a-glance
# scanning -- overridden by _INVALID_FILL when that same cell is also
# flagged, since a flagged value being wrong matters more than what it
# currently says.
_STATUS_FILLS = {
    "Present": PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid"),
    "Weekoff": PatternFill(start_color="C9DAF8", end_color="C9DAF8", fill_type="solid"),
    "LEAVE": PatternFill(start_color="FCE5CD", end_color="FCE5CD", fill_type="solid"),
}


def generate_excel_report(data, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Report"

    _write_header(ws)
    row = 2

    if "employees" in data:
        for employee_data in data.get("employees", []):
            row = _write_employee_rows(ws, row, employee_data)
    else:
        row = _write_employee_rows(ws, row, data)

    for col_idx, width in enumerate(_COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    wb.save(output_path)
    return output_path


def _write_header(ws):
    for col_idx, title in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _BORDER


def _write_employee_rows(ws, row, employee_data):
    code = employee_data.get("employee_code") or ""
    name = employee_data.get("employee_name") or "Not Detected"

    for record in employee_data.get("records", []):
        cell_highlight = record.get("cell_highlight", [])
        row_fully_invalid = len(cell_highlight) == 4
        status = record.get("status", "")

        values = [
            code,
            name,
            record.get("date", ""),
            _SHIFT_CODE,
            _SHIFT_TIME,
            (record.get("in_time", "") or "").strip(),
            (record.get("out_time", "") or "").strip(),
            status,
            (record.get("remark", "") or "").strip(),
        ]

        for col_idx, (title, value) in enumerate(zip(_COLUMNS, values), start=1):
            cell = ws.cell(row=row, column=col_idx, value=value)
            cell.alignment = _CENTER
            cell.border = _BORDER

            field = next((f for f, c in _FIELD_TO_COLUMN.items() if c == title), None)
            if row_fully_invalid or (field and field in cell_highlight):
                cell.fill = _INVALID_FILL
                cell.font = Font(bold=True)
            elif title == "Status" and status in _STATUS_FILLS:
                cell.fill = _STATUS_FILLS[status]

        row += 1

    return row
